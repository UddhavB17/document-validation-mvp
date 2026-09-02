"""Report generation for JSON and Excel anomaly reports."""

import json
import re
from datetime import datetime, timezone
from pathlib import Path

from database.db import get_connection
from services.paths import report_output_dir
from services.processing_policy import is_internal_document_type

REPORTS_DIR = report_output_dir()
REPORT_DIR = REPORTS_DIR


def generate_excel_report(application_id: int) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except Exception as exc:
        raise RuntimeError("openpyxl is required to generate Excel reports") from exc

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    data = _load_report_data(application_id)
    application = data["application"]
    anomalies = data["anomalies"]
    uploaded_file = data["uploaded_file"]

    loan_id = application.get("loan_id") or f"application_{application_id}"
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    output_path = REPORT_DIR / f"{_safe_filename_part(str(loan_id))}_{timestamp}.xlsx"

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    anomaly_sheet = workbook.create_sheet("Anomalies")

    summary_rows = [
        ("Loan ID", application.get("loan_id")),
        ("Applicant Name", application.get("applicant_name")),
        ("Product Type", application.get("product_type")),
        ("Branch", application.get("branch")),
        ("Total Pages", uploaded_file.get("total_pages")),
        ("Digital Pages", uploaded_file.get("digital_pages")),
        ("Scanned Pages", uploaded_file.get("scanned_pages")),
        ("Documents Found", ", ".join(data["documents_found"])),
        ("Issues Count", len(anomalies)),
        ("Final Status", application.get("status")),
        ("Generated At", datetime.now().isoformat(timespec="seconds")),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=1):
        summary_sheet.cell(row=row_index, column=1, value=label)
        summary_sheet.cell(row=row_index, column=2, value=value)

    anomaly_headers = ["Rule ID", "Severity", "Document", "Expected", "Found", "Page", "Reason"]
    anomaly_sheet.append(anomaly_headers)
    for anomaly in anomalies:
        anomaly_sheet.append(
            [
                anomaly.get("rule_id"),
                anomaly.get("severity"),
                anomaly.get("document_type"),
                anomaly.get("expected_value"),
                anomaly.get("found_value"),
                anomaly.get("page_number"),
                anomaly.get("reason"),
            ]
        )

    severity_fills = {
        "HIGH": PatternFill(fill_type="solid", fgColor="FF4444"),
        "MEDIUM": PatternFill(fill_type="solid", fgColor="FFA500"),
        "LOW": PatternFill(fill_type="solid", fgColor="FFFF00"),
    }
    for sheet in [summary_sheet, anomaly_sheet]:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        _auto_size_columns(sheet)

    for row in anomaly_sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        cell = row[0]
        fill = severity_fills.get(str(cell.value).upper())
        if fill:
            cell.fill = fill

    workbook.save(output_path)
    return str(output_path)


def build_report(
    application_id: int,
    loan_id: str | None = None,
    exceptions: list[dict] | None = None,
    llm_summary: str = "",
    metadata: dict | None = None,
) -> dict:
    exceptions = exceptions or []
    from services.reviewer import collapse_for_reviewer

    actionable = collapse_for_reviewer(exceptions)
    report = {
        "application_id": application_id,
        "loan_id": loan_id,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "total_exceptions": len(exceptions),
        "actionable_exception_count": len(actionable),
        "high_severity_count": sum(1 for e in actionable if str(e.get("severity")).upper() == "HIGH"),
        "llm_summary": llm_summary,
        "exceptions": exceptions,
        "actionable_exceptions": actionable,
    }
    if metadata:
        report["metadata"] = metadata
    return report


def save_report_json(report: dict) -> Path:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    loan_id = _safe_filename_part(str(report.get("loan_id") or "unknown"))
    filename = f"report_{report['application_id']}_{loan_id}.json"
    output_path = REPORTS_DIR / filename
    with output_path.open("w", encoding="utf-8") as file:
        json.dump(report, file, indent=2, ensure_ascii=False)
    return output_path


def _safe_filename_part(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "_", value.strip())
    cleaned = cleaned.strip("._")
    return cleaned or "unknown"


def _load_report_data(application_id: int) -> dict:
    with get_connection() as connection:
        application = connection.execute(
            "SELECT * FROM applications WHERE id = ?",
            (application_id,),
        ).fetchone()
        if application is None:
            raise ValueError(f"Application {application_id} not found")

        uploaded_file = connection.execute(
            """
            SELECT * FROM uploaded_files
            WHERE application_id = ?
            ORDER BY uploaded_at DESC
            LIMIT 1
            """,
            (application_id,),
        ).fetchone()
        anomalies = connection.execute(
            """
            SELECT rule_id, s_no, severity, document_type, expected_value, found_value, page_number, reason
            FROM validation_results
            WHERE application_id = ?
            ORDER BY
              CASE severity WHEN 'HIGH' THEN 1 WHEN 'MEDIUM' THEN 2 WHEN 'LOW' THEN 3 ELSE 4 END,
              page_number
            """,
            (application_id,),
        ).fetchall()
        pages = connection.execute(
            """
            SELECT DISTINCT document_type
            FROM pages
            WHERE application_id = ? AND document_type IS NOT NULL AND document_type != 'Unknown'
            """,
            (application_id,),
        ).fetchall()

    return {
        "application": dict(application),
        "uploaded_file": dict(uploaded_file) if uploaded_file else {},
        "anomalies": [dict(row) for row in anomalies],
        "documents_found": sorted(
            row["document_type"]
            for row in pages
            if not is_internal_document_type(row["document_type"])
        ),
    }


def _auto_size_columns(sheet) -> None:
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 60)
