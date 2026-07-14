from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
import pytest
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from services.zip_package import PackageValidationError, normalize_zip_package


def _create_pdf(path: Path, pages: int) -> None:
    document = fitz.open()
    for page_number in range(1, pages + 1):
        page = document.new_page()
        page.insert_text((72, 72), f"Source page {page_number}")
    document.save(path)
    document.close()


def test_zip_package_normalizes_documents_and_preserves_source_ranges(tmp_path: Path) -> None:
    first_pdf = tmp_path / "first.pdf"
    second_pdf = tmp_path / "second.pdf"
    image_path = tmp_path / "aadhaar.jpg"
    _create_pdf(first_pdf, 2)
    _create_pdf(second_pdf, 1)
    Image.new("RGB", (100, 60), "white").save(image_path)

    archive_path = tmp_path / "loan.zip"
    with ZipFile(archive_path, "w", ZIP_DEFLATED) as archive:
        archive.write(first_pdf, "Applicant/PAN.pdf")
        archive.writestr("__MACOSX/._PAN.pdf", b"metadata")
        archive.write(image_path, "Applicant/Aadhaar.jpg")
        archive.write(second_pdf, "CoApplicant/PAN.pdf")

    result = normalize_zip_package(archive_path, tmp_path / "package")

    assert result["total_files"] == 3
    assert result["total_pages"] == 4
    assert [item["source_document_id"] for item in result["documents"]] == [
        "file-0001",
        "file-0002",
        "file-0003",
    ]
    assert [item["pages"] for item in result["documents"]] == [[1, 2], [3], [4]]
    normalized = fitz.open(result["normalized_pdf_path"])
    try:
        assert normalized.page_count == 4
    finally:
        normalized.close()


@pytest.mark.parametrize(
    ("member_name", "message"),
    [
        ("../outside.pdf", "Unsafe path"),
        ("nested/documents.zip", "Nested ZIP"),
        ("notes.txt", "Unsupported files"),
    ],
)
def test_zip_package_rejects_unsafe_or_unsupported_members(
    tmp_path: Path,
    member_name: str,
    message: str,
) -> None:
    archive_path = tmp_path / "invalid.zip"
    with ZipFile(archive_path, "w") as archive:
        archive.writestr(member_name, b"not a supported document")

    with pytest.raises(PackageValidationError, match=message):
        normalize_zip_package(archive_path, tmp_path / "package")


def test_zip_package_rejects_corrupt_pdf(tmp_path: Path) -> None:
    archive_path = tmp_path / "corrupt.zip"
    with ZipFile(archive_path, "w") as archive:
        archive.writestr("fake.pdf", b"this is not a PDF")

    with pytest.raises(PackageValidationError, match="Unreadable document"):
        normalize_zip_package(archive_path, tmp_path / "package")


def test_zip_package_renders_nonempty_xlsx_worksheets_to_internal_pages(tmp_path: Path) -> None:
    workbook_path = tmp_path / "bank-statement.xlsx"
    workbook = Workbook()
    statement = workbook.active
    statement.title = "Transactions"
    statement.append(["Date", "Description", "Debit", "Credit", "Balance"])
    statement.append(["2026-07-01", "Opening balance", "", "1000", "1000"])
    # Excel formatting can inflate max_row even when there is no data there.
    statement["A50001"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    workbook.create_sheet("Empty Sheet")
    workbook.save(workbook_path)
    workbook.close()

    archive_path = tmp_path / "loan.zip"
    with ZipFile(archive_path, "w", ZIP_DEFLATED) as archive:
        archive.write(workbook_path, "Bank/bank-statement.xlsx")

    result = normalize_zip_package(archive_path, tmp_path / "package")

    assert result["total_files"] == 1
    assert result["total_pages"] == 1
    assert result["documents"][0]["file_type"] == "xlsx"
    assert result["documents"][0]["worksheets"] == ["Transactions"]
    normalized = fitz.open(result["normalized_pdf_path"])
    try:
        text = normalized[0].get_text()
    finally:
        normalized.close()
    assert "Transactions" in text
    assert "Opening balance" in text


def test_zip_package_limits_actual_xlsx_values_not_formatting_bounds(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook_path = tmp_path / "too-many-values.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Header"])
    worksheet.append(["Value 1"])
    worksheet.append(["Value 2"])
    workbook.save(workbook_path)
    workbook.close()
    archive_path = tmp_path / "loan.zip"
    with ZipFile(archive_path, "w", ZIP_DEFLATED) as archive:
        archive.write(workbook_path, "too-many-values.xlsx")
    monkeypatch.setenv("MAX_XLSX_ROWS_PER_SHEET", "2")

    with pytest.raises(PackageValidationError, match="non-empty rows"):
        normalize_zip_package(archive_path, tmp_path / "package")
